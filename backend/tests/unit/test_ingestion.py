from __future__ import annotations

from io import BytesIO

import pytest

from services.ingestion import IngestionService


def test_csv_is_classified_as_transaction_data() -> None:
    result = IngestionService().inspect(
        "transactions.csv",
        b"date,account,account_type,amount\n2026-08-01,Sales,revenue,1200.00\n",
        "text/csv",
    )

    assert result.classification == "transaction_data"
    assert result.tables[0].rows[0]["amount"] == "1200.00"
    assert result.review_required is False


def test_period_summary_headers_do_not_route_to_transaction_parser() -> None:
    result = IngestionService().inspect(
        "period_summaries.csv",
        b"period_id,start_date,end_date,account,account_type,amount\n2026-08,2026-08-01,2026-08-31,Sales,revenue,1200\n",
        "text/csv",
    )

    assert result.classification == "period_summary"


def test_xlsx_extracts_a_reviewable_financial_table() -> None:
    openpyxl = pytest.importorskip("openpyxl")
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "August"
    sheet.append(["date", "account", "account_type", "amount"])
    sheet.append(["2026-08-01", "Sales", "revenue", 1200])
    buffer = BytesIO()
    workbook.save(buffer)

    result = IngestionService().inspect("august.xlsx", buffer.getvalue())

    assert result.classification == "transaction_data"
    assert result.review_required is False
    assert result.tables[0].source.sheet == "August"


def test_reference_text_never_becomes_financial_data() -> None:
    result = IngestionService().inspect("notes.txt", b"Renewal discussion: customer expects a Q4 expansion.")

    assert result.classification == "reference_document"
    assert not result.tables


def test_docx_table_is_extracted_with_table_provenance() -> None:
    from docx import Document

    document = Document()
    table = document.add_table(rows=2, cols=4)
    for index, value in enumerate(["date", "account", "account_type", "amount"]):
        table.rows[0].cells[index].text = value
    for index, value in enumerate(["2026-08-01", "Sales", "revenue", "1200"]):
        table.rows[1].cells[index].text = value
    buffer = BytesIO()
    document.save(buffer)

    result = IngestionService().inspect("statement.docx", buffer.getvalue())

    assert result.classification == "transaction_data"
    assert result.tables[0].source.table == 1


def test_native_pdf_text_is_reference_only_without_ocr() -> None:
    fitz = pytest.importorskip("fitz")
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "August management note: enterprise accounts renewed on schedule.")

    result = IngestionService().inspect("notes.pdf", document.tobytes())

    assert result.classification == "reference_document"
    assert result.requires_ocr is False
    assert result.page_count == 1


def test_unsupported_format_has_an_actionable_error() -> None:
    with pytest.raises(ValueError, match="Unsupported format"):
        IngestionService().inspect("macro.xlsm", b"not really an Excel file")
