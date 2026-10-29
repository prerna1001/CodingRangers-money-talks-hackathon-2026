"""Pitch deck export (GET /api/runs/{run_id}/pitch-deck).

Runs the demo pipeline, downloads both themed decks and checks the shape a
reader relies on: 5-7 slides, a native chart on the bridge and driver slides,
speaker notes throughout, and a clean 4xx for a bad theme or unknown run.
"""

from __future__ import annotations

import io
import json
import re
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE

from api.routes.analyze import _RUNS
from graph.workflow import run_pipeline
from main import app
from config import DEFAULT_USER_ID
from models.schemas import CanonicalTransaction, CompanyProfileCore, PeriodSummary
from services.pitch_deck import PALETTES

FIXTURES_DIR = Path(__file__).resolve().parent.parent / "fixtures"


@pytest.fixture(autouse=True)
def isolated_memory_db(tmp_path, monkeypatch):
    monkeypatch.setattr("services.memory_store._DEFAULT_DB_PATH", str(tmp_path / "test_memory.db"))


@pytest.fixture()
def finished_run():
    data = json.loads((FIXTURES_DIR / "saas_2026_08_transactions.json").read_text())
    state = run_pipeline(
        company_id="demo_saas_test",
        current_period="2026-08",
        prior_period="2026-07",
        company_profile_core=CompanyProfileCore(
            company_id="demo_saas_test", company_name="DemoCo", industry="B2B SaaS", business_model="subscription"
        ),
        transactions=[CanonicalTransaction(**t) for t in data["transactions"]],
        period_summaries=[PeriodSummary(**p) for p in data["period_summaries"]],
    )
    run_id = state["run_id"]
    _RUNS[run_id] = {"status": "complete", "state": state, "error": None, "user_id": DEFAULT_USER_ID}
    yield run_id
    _RUNS.pop(run_id, None)


@pytest.mark.parametrize("theme", list(PALETTES))
def test_deck_has_pitch_structure(finished_run, theme):
    response = TestClient(app).get(f"/api/runs/{finished_run}/pitch-deck", params={"theme": theme})
    assert response.status_code == 200, response.text
    assert response.headers["content-disposition"].endswith(f'_pitch_{theme}.pptx"')
    slides = list(Presentation(io.BytesIO(response.content)).slides)
    assert 5 <= len(slides) <= 7
    charts = [sum(1 for shape in slide.shapes if shape.has_chart) for slide in slides]
    assert charts[2] >= 1 and charts[3] >= 1  # bridge and drivers are native charts
    assert all(slide.has_notes_slide and slide.notes_slide.notes_text_frame.text for slide in slides)


def test_unknown_theme_and_run_are_rejected(finished_run):
    client = TestClient(app)
    assert client.get(f"/api/runs/{finished_run}/pitch-deck", params={"theme": "neon"}).status_code == 422
    assert client.get("/api/runs/run_missing/pitch-deck").status_code == 404


def test_slides_address_review_comments(finished_run):
    response = TestClient(app).get(f"/api/runs/{finished_run}/pitch-deck")
    slides = list(Presentation(io.BytesIO(response.content)).slides)

    def texts(slide):
        return " ".join(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame)

    # Slide 3: waterfall with % labels plus key / non-key drivers, increases vs decreases.
    bridge = next(shape.chart for shape in slides[2].shapes if shape.has_chart)
    assert any("%)" in name for name in bridge.plots[0].categories)
    assert '"K"' in bridge.value_axis.tick_labels.number_format
    assert "INCREASES" in texts(slides[2]) and "DECREASES" in texts(slides[2]) and "KEY" in texts(slides[2])

    # Slide 4: graphs instead of the table, including a pie with percentages.
    assert not any(shape.has_table for shape in slides[3].shapes)
    pie = next(shape.chart for shape in slides[3].shapes
               if shape.has_chart and shape.chart.chart_type == XL_CHART_TYPE.PIE)
    assert pie.plots[0].data_labels.show_percentage

    # Slide 6: three columns, one of them recommendations.
    for heading in ("Key findings", "Risks and caveats", "Recommendations"):
        assert heading in texts(slides[5])

    # Thousands, never millions.
    assert not re.search(r"\$[\d.,]+M\b", " ".join(texts(slide) for slide in slides))
