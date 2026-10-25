"""Power BI package export (GET /api/runs/{run_id}/powerbi).

Runs the demo pipeline in LLM_MOCK_MODE, then checks the single-workbook
package: every sheet present, evidence only points at real facts and
transactions, the bridge is waterfall-ready without losing the closing
total, only the two join keys are shared between sheets (so Power BI's
relationship autodetect cannot guess wrong), and formula-looking text stays
text.
"""

from __future__ import annotations

import base64
import io
import json
import re
import zipfile
import zlib
from collections import Counter
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from api.routes.analyze import _RUNS
from graph.workflow import run_pipeline
from main import app
from models.schemas import CanonicalTransaction, CompanyProfileCore, PeriodSummary
from services.powerbi_export import SHEETS
from services.powerbi_report import THEMES as REPORT_THEMES

FIXTURES_DIR = Path(__file__).resolve().parent.parent / "fixtures"
INJECTED_MEMO = '=HYPERLINK("http://example.test","click")'


@pytest.fixture(autouse=True)
def isolated_memory_db(tmp_path, monkeypatch):
    monkeypatch.setattr("services.memory_store._DEFAULT_DB_PATH", str(tmp_path / "test_memory.db"))


@pytest.fixture()
def finished_run():
    data = json.loads((FIXTURES_DIR / "saas_2026_08_transactions.json").read_text())
    transactions = [CanonicalTransaction(**t) for t in data["transactions"]]
    transactions[0] = transactions[0].model_copy(update={"memo": INJECTED_MEMO})
    state = run_pipeline(
        company_id="demo_saas_test",
        current_period="2026-08",
        prior_period="2026-07",
        company_profile_core=CompanyProfileCore(
            company_id="demo_saas_test",
            company_name="DemoCo",
            industry="B2B SaaS",
            business_model="subscription",
        ),
        transactions=transactions,
        period_summaries=[PeriodSummary(**p) for p in data["period_summaries"]],
    )
    run_id = state["run_id"]
    _RUNS[run_id] = {"status": "complete", "state": state, "error": None}
    yield run_id
    _RUNS.pop(run_id, None)


def _sheet_rows(workbook, name: str) -> list[dict]:
    rows = list(workbook[name].iter_rows(values_only=True))
    header = rows[0]
    return [dict(zip(header, row)) for row in rows[1:]]


def _package(run_id: str) -> zipfile.ZipFile:
    response = TestClient(app).get(f"/api/runs/{run_id}/powerbi")
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == "application/zip"
    return zipfile.ZipFile(io.BytesIO(response.content))


def test_package_contains_reports_data_and_themes(finished_run):
    names = set(_package(finished_run).namelist())
    expected = {
        "README.txt",
        f"data/Ledgerlight_{finished_run}.xlsx",
        "reports/Ledgerlight.SemanticModel/model.bim",
        "reports/Ledgerlight.SemanticModel/definition.pbism",
        "themes/Ledgerlight App.json",
    }
    for theme_key, theme in REPORT_THEMES.items():
        expected |= {
            f"reports/Ledgerlight {theme_key}.pbip",
            f"reports/Ledgerlight {theme_key}.Report/definition.pbir",
            f"reports/Ledgerlight {theme_key}.Report/report.json",
            f"themes/{theme['name']}.json",
        }
    for table in SHEETS:
        expected |= {f"data/csv/{table}.csv", f"data/powerquery/{table}.pq"}
    assert expected <= names


def test_pbip_projects_point_at_files_that_exist(finished_run):
    archive = _package(finished_run)
    names = set(archive.namelist())
    for theme_key in REPORT_THEMES:
        pbip = json.loads(archive.read(f"reports/Ledgerlight {theme_key}.pbip"))
        report_folder = pbip["artifacts"][0]["report"]["path"]
        pbir = json.loads(archive.read(f"reports/{report_folder}/definition.pbir"))
        model_folder = pbir["datasetReference"]["byPath"]["path"].removeprefix("../")
        assert f"reports/{model_folder}/model.bim" in names
        report = json.loads(archive.read(f"reports/{report_folder}/report.json"))
        for item in report["resourcePackages"][0]["resourcePackage"]["items"]:
            assert f"reports/{report_folder}/StaticResources/RegisteredResources/{item['path']}" in names


def test_semantic_model_embeds_every_row(finished_run):
    archive = _package(finished_run)
    model = json.loads(archive.read("reports/Ledgerlight.SemanticModel/model.bim"))["model"]
    workbook = load_workbook(io.BytesIO(archive.read(f"data/Ledgerlight_{finished_run}.xlsx")))
    for table in model["tables"]:
        expression = "\n".join(table["partitions"][0]["source"]["expression"])
        packed = re.search(r'Binary\.FromText\("([^"]+)"', expression).group(1)
        rows = json.loads(zlib.decompress(base64.b64decode(packed), -15))
        assert len(rows) == workbook[table["name"]].max_row - 1, table["name"]
        assert [c["name"] for c in table["columns"]] == list(SHEETS[table["name"]])
    keys = {(r["fromTable"], r["fromColumn"], r["toTable"], r["toColumn"]) for r in model["relationships"]}
    assert keys == {("Evidence", "fact_id", "Facts", "fact_id"), ("Evidence", "txn_id", "Transactions", "txn_id")}


def test_every_visual_field_exists_in_the_model(finished_run):
    archive = _package(finished_run)
    model = json.loads(archive.read("reports/Ledgerlight.SemanticModel/model.bim"))["model"]
    columns = {(t["name"], c["name"]) for t in model["tables"] for c in t["columns"]}
    measures = {(t["name"], m["name"]) for t in model["tables"] for m in t.get("measures", [])}
    for theme_key in REPORT_THEMES:
        report = json.loads(archive.read(f"reports/Ledgerlight {theme_key}.Report/report.json"))
        assert [s["displayName"] for s in report["sections"]] == ["Overview", "Evidence", "Trend and transactions"]
        for section in report["sections"]:
            for container in section["visualContainers"]:
                single = json.loads(container["config"])["singleVisual"]
                query = single.get("prototypeQuery")
                if not query:
                    continue
                entities = {f["Name"]: f["Entity"] for f in query["From"]}
                for select in query["Select"]:
                    node = select.get("Aggregation", {}).get("Expression", select)
                    if "Measure" in node:
                        ref = node["Measure"]
                        assert (entities[ref["Expression"]["SourceRef"]["Source"]], ref["Property"]) in measures
                    else:
                        ref = node["Column"]
                        assert (entities[ref["Expression"]["SourceRef"]["Source"]], ref["Property"]) in columns


def test_overview_has_percent_pie_and_thousands(finished_run):
    archive = _package(finished_run)
    for theme_key in REPORT_THEMES:
        report = json.loads(archive.read(f"reports/Ledgerlight {theme_key}.Report/report.json"))
        overview = [json.loads(c["config"])["singleVisual"] for c in report["sections"][0]["visualContainers"]]
        pie = next(v for v in overview if v["visualType"] == "pieChart")
        assert "percent of total" in pie["objects"]["labels"][0]["properties"]["labelStyle"]["expr"]["Literal"]["Value"]
        for visual in overview:
            if visual["visualType"] in ("columnChart", "clusteredBarChart"):
                units = visual["objects"]["labels"][0]["properties"]["labelDisplayUnits"]
                assert units["expr"]["Literal"]["Value"] == "1000D"
        # The bridge is stacked columns whose carried-forward series fills the
        # space under each driver bar; no floating bars over empty space.
        bridge_visual = next(v for v in overview if v["visualType"] == "columnChart")
        fills = {d["selector"]["metadata"] for d in bridge_visual["objects"]["dataPoint"]}
        assert "Sum(BridgeChart.chart_carried)" in fills

    workbook = load_workbook(io.BytesIO(archive.read(f"data/Ledgerlight_{finished_run}.xlsx")))
    bridge = _sheet_rows(workbook, "Bridge")
    assert all("%)" in row["bridge_label"] for row in bridge if row["bridge_type"] != "start")
    chart = _sheet_rows(workbook, "BridgeChart")
    summary = _sheet_rows(workbook, "Summary")[0]
    assert chart[0]["chart_total"] == summary["from_revenue"] and chart[-1]["chart_total"] == summary["to_revenue"]
    for row in chart[1:-1]:
        # Each driver column stands on the running total, which ties to the close.
        assert row["chart_carried"] is not None
    assert round(chart[-2]["chart_carried"] + (chart[-2]["chart_increase"] or 0), 2) == round(summary["to_revenue"], 2)
    assert all(row["driver_label"].endswith(("%)", "(new)")) for row in _sheet_rows(workbook, "Drivers"))


def test_powerbi_package_is_one_ready_to_import_workbook(finished_run):
    archive = _package(finished_run)
    workbook_name = f"data/Ledgerlight_{finished_run}.xlsx"

    workbook = load_workbook(io.BytesIO(archive.read(workbook_name)))
    assert workbook.sheetnames == list(SHEETS)
    for name, columns in SHEETS.items():
        assert [c.value for c in workbook[name][1]] == list(columns)

    transactions = _sheet_rows(workbook, "Transactions")
    facts = _sheet_rows(workbook, "Facts")
    evidence = _sheet_rows(workbook, "Evidence")
    assert transactions and facts and evidence
    txn_ids = {row["txn_id"] for row in transactions}
    fact_ids = {row["fact_id"] for row in facts}
    assert all(row["txn_id"] in txn_ids and row["fact_id"] in fact_ids for row in evidence)

    # Waterfall-ready: numbered bars, no closing bar, and the bars plus
    # Power BI's own Total equal the closing revenue kept in Summary.
    summary = _sheet_rows(workbook, "Summary")[0]
    bridge = _sheet_rows(workbook, "Bridge")
    assert bridge[0]["bridge_type"] == "start"
    assert all(row["bridge_type"] != "end" for row in bridge)
    assert [row["bar"] for row in bridge] == sorted(row["bar"] for row in bridge)
    assert round(sum(row["bridge_amount"] for row in bridge), 2) == round(summary["to_revenue"], 2)


def test_only_join_keys_are_shared_between_sheets():
    counts = Counter(column for columns in SHEETS.values() for column in columns)
    shared = {column for column, count in counts.items() if count > 1}
    assert shared == {"fact_id", "txn_id"}


def test_formula_like_memo_stays_text(finished_run):
    archive = _package(finished_run)
    workbook = load_workbook(io.BytesIO(archive.read(f"data/Ledgerlight_{finished_run}.xlsx")))
    memo_column = list(SHEETS["Transactions"]).index("memo")
    cells = [row[memo_column] for row in workbook["Transactions"].iter_rows(min_row=2)]
    injected = [cell for cell in cells if cell.value == INJECTED_MEMO]
    assert injected and all(cell.data_type == "s" for cell in injected)


def test_powerbi_package_404s_for_unknown_run():
    assert TestClient(app).get("/api/runs/run_missing/powerbi").status_code == 404
