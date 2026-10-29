"""Pitch deck for one run, as a native PowerPoint file
(GET /api/runs/{run_id}/pitch-deck?theme=finance|dark).

Seven 16:9 slides written for a reader who was not in the room -- another
team, a client, senior leadership:

    1  Title and headline result
    2  Problem and approach
    3  Revenue bridge            (waterfall + key drivers, increases vs decreases)
    4  Driver analysis           (driver bar chart + key vs non-key pie)

Money is shown in thousands ($183K, never $0.18M) and chart labels carry
percentages, matching the Power BI report.
    5  Trust and evidence        (quality metrics + agent run chart)
    6  Findings, risks, recommendations (three columns)
    7  Future work and next steps

Charts are native PowerPoint charts, not pictures, so they stay editable
and can be restyled in PowerPoint. The two themes mirror the two Power BI
reports (services/powerbi_report.py): corporate navy and executive dark.
Everything shown comes from the same run state the Power BI export reads;
nothing is invented, so a number on a slide always matches the dashboard.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from typing import Any

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from api.presenters import to_frontend_analysis
from models.schemas import RunState

FONT = "Calibri"
FONT_BOLD = "Calibri"
SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.6)
CONTENT_W = SLIDE_W - 2 * MARGIN

# ---------------------------------------------------------------------------
# Themes -- same palettes as the two Power BI reports
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class Palette:
    key: str
    label: str
    page: str
    card: str
    border: str
    text: str
    muted: str
    accent: str       # titles, totals, header fills
    accent_2: str
    good: str
    bad: str
    on_accent: str
    series: tuple[str, ...]
    slices: tuple[str, ...]  # pie slices for key drivers: no greys (non-key) or reds (losses)
    carry: str               # waterfall: revenue carried forward under each driver bar


PALETTES = {
    "finance": Palette(
        key="finance", label="Corporate Navy", page="F3F5F8", card="FFFFFF", border="DDE3EA",
        text="1F2A44", muted="5B6778", accent="1F3864", accent_2="2E75B6", good="2E7D32", bad="C62828",
        on_accent="FFFFFF", series=("1F3864", "2E75B6", "00A3A1", "F2A900", "7F8C9D", "C0504D"),
        slices=("1F3864", "2E75B6", "00A3A1", "F2A900", "70AD47", "8064A2"), carry="D3DBE6",
    ),
    "dark": Palette(
        key="dark", label="Executive Dark", page="0E1621", card="16212E", border="26364A",
        text="E6EDF5", muted="9AA8B8", accent="5B9BD5", accent_2="2BB3A3", good="3FB97A", bad="E5534B",
        on_accent="0E1621", series=("5B9BD5", "2BB3A3", "F2C14E", "A5B4C8", "E86A5F", "7CC47F"),
        slices=("5B9BD5", "2BB3A3", "F2C14E", "B18CE0", "7CC47F", "4E79A7"), carry="2B3B50",
    ),
}


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


# ---------------------------------------------------------------------------
# Formatting
# ---------------------------------------------------------------------------


# Chart labels and axes in thousands ($183K), matching the Power BI report.
# Excel scales by 1,000 per trailing comma in a number format.
CHART_K = '[>=10000]$#,##0,"K";$#,##0.0,"K"'
AXIS_K = '$#,##0,"K"'


def _money(value: float, signed: bool = False) -> str:
    """Thousands, never millions: $1,250K rather than $1.25M."""
    sign = "-" if value < 0 else ("+" if signed else "")
    v = abs(value)
    if v >= 10_000:
        body = f"${v / 1_000:,.0f}K"
    elif v >= 1_000:
        body = f"${v / 1_000:,.1f}K"
    else:
        body = f"${v:,.0f}"
    return sign + body


def _pct(value: float, signed: bool = False) -> str:
    return f"{value:+.1f}%" if signed else f"{value:.1f}%"


def _clip(text: str, limit: int) -> str:
    text = " ".join(str(text).split())
    return text if len(text) <= limit else text[: limit - 1].rstrip() + "…"


def _first_sentence(text: str, limit: int = 220) -> str:
    text = " ".join((text or "").split())
    end = text.find(". ")
    if 0 < end < limit:
        return text[: end + 1]
    return _clip(text, limit)


# ---------------------------------------------------------------------------
# Drawing helpers
# ---------------------------------------------------------------------------


class _Deck:
    def __init__(self, palette: Palette):
        self.p = palette
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = SLIDE_W, SLIDE_H
        self.blank = self.prs.slide_layouts[6]
        self.count = 0

    # -- primitives ---------------------------------------------------------

    def rect(self, slide, x, y, w, h, fill: str | None, line: str | None = None, shape=MSO_SHAPE.RECTANGLE):
        shp = slide.shapes.add_shape(shape, x, y, w, h)
        shp.shadow.inherit = False
        if fill:
            shp.fill.solid()
            shp.fill.fore_color.rgb = _rgb(fill)
        else:
            shp.fill.background()
        if line:
            shp.line.color.rgb = _rgb(line)
            shp.line.width = Pt(0.75)
        else:
            shp.line.fill.background()
        return shp

    def text(self, slide, x, y, w, h, paragraphs, *, size=14, color=None, bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, space_after=0, line_spacing=1.0):
        """paragraphs: str | list of str | list of (str, {size,bold,color,bullet}) tuples."""
        box = slide.shapes.add_textbox(x, y, w, h)
        frame = box.text_frame
        frame.word_wrap = True
        frame.vertical_anchor = anchor
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        if isinstance(paragraphs, str):
            paragraphs = [paragraphs]
        for index, item in enumerate(paragraphs):
            content, opts = (item, {}) if isinstance(item, str) else item
            para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            para.alignment = align
            para.space_after = Pt(opts.get("space_after", space_after))
            para.line_spacing = line_spacing
            run = para.add_run()
            run.text = content
            font = run.font
            font.name = FONT_BOLD if opts.get("bold", bold) else FONT
            font.size = Pt(opts.get("size", size))
            font.bold = opts.get("bold", bold)
            font.color.rgb = _rgb(opts.get("color", color or self.p.text))
            if opts.get("bullet"):
                pPr = para._p.get_or_add_pPr()
                pPr.set("marL", str(Inches(0.2)))
                pPr.set("indent", str(-Inches(0.2)))
                for tag in ("a:buNone", "a:buChar"):
                    for el in pPr.findall(qn(tag)):
                        pPr.remove(el)
                bu_clr = pPr.makeelement(qn("a:buClr"), {})
                srgb = bu_clr.makeelement(qn("a:srgbClr"), {"val": self.p.accent_2})
                bu_clr.append(srgb)
                pPr.append(bu_clr)
                pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))
        return box

    def notes(self, slide, text: str):
        slide.notes_slide.notes_text_frame.text = text

    # -- slide scaffolding --------------------------------------------------

    def slide(self, section: str, title: str, run_label: str):
        s = self.prs.slides.add_slide(self.blank)
        self.count += 1
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = _rgb(self.p.page)
        # Left accent rail + section label + action title
        self.rect(s, 0, 0, Inches(0.18), SLIDE_H, self.p.accent)
        self.text(s, MARGIN, Inches(0.38), CONTENT_W, Inches(0.3), section.upper(), size=11, bold=True,
                  color=self.p.accent_2)
        self.text(s, MARGIN, Inches(0.68), CONTENT_W, Inches(0.95), title, size=26, bold=True,
                  color=self.p.accent if self.p.key == "finance" else self.p.text, line_spacing=0.95)
        # Footer
        self.rect(s, MARGIN, Inches(7.0), CONTENT_W, Emu(9525), self.p.border)
        self.text(s, MARGIN, Inches(7.07), Inches(9), Inches(0.25), run_label, size=9, color=self.p.muted)
        self.text(s, SLIDE_W - MARGIN - Inches(1), Inches(7.07), Inches(1), Inches(0.25), str(self.count), size=9,
                  color=self.p.muted, align=PP_ALIGN.RIGHT)
        return s

    def card(self, slide, x, y, w, h):
        return self.rect(slide, x, y, w, h, self.p.card, self.p.border)

    def kpi(self, slide, x, y, w, h, value, label, tone: str | None = None, size=26):
        self.card(slide, x, y, w, h)
        self.rect(slide, x, y, Inches(0.06), h, tone or self.p.accent_2)
        self.text(slide, x + Inches(0.25), y + Inches(0.14), w - Inches(0.4), Inches(0.3), label.upper(), size=10,
                  bold=True, color=self.p.muted)
        self.text(slide, x + Inches(0.25), y + Inches(0.46), w - Inches(0.4), h - Inches(0.5), value, size=size,
                  bold=True, color=tone or self.p.text, anchor=MSO_ANCHOR.TOP)

    def panel_title(self, slide, x, y, w, label):
        self.text(slide, x, y, w, Inches(0.3), label, size=13, bold=True, color=self.p.text)
        self.rect(slide, x, y + Inches(0.32), Inches(0.5), Emu(25400), self.p.accent_2)

    # -- chart styling ------------------------------------------------------

    def style_chart(self, chart, *, legend=False):
        chart.font.name = FONT
        chart.font.size = Pt(10)
        chart.font.color.rgb = _rgb(self.p.muted)
        chart.has_legend = legend
        chart.has_title = False

    def style_axes(self, chart, *, value_format=AXIS_K, gridlines=True):
        cat, val = chart.category_axis, chart.value_axis
        cat.format.line.color.rgb = _rgb(self.p.border)
        cat.tick_labels.font.size = Pt(10)
        cat.tick_labels.font.color.rgb = _rgb(self.p.text)
        cat.has_major_gridlines = False
        val.has_major_gridlines = gridlines
        if gridlines:
            val.major_gridlines.format.line.color.rgb = _rgb(self.p.border)
            val.major_gridlines.format.line.width = Pt(0.5)
        val.format.line.fill.background()
        val.tick_labels.number_format = value_format
        val.tick_labels.number_format_is_linked = False
        val.tick_labels.font.size = Pt(9)
        val.tick_labels.font.color.rgb = _rgb(self.p.muted)

    def table(self, slide, x, y, w, rows: list[list[str]], widths: list[float], *, row_h=0.36, align_right_from=1):
        shape = slide.shapes.add_table(len(rows), len(rows[0]), x, y, w, Inches(row_h * len(rows)))
        table = shape.table
        # Strip the default style so the cells carry the theme.
        tbl_pr = shape._element.graphic.graphicData.tbl.tblPr
        for attr in ("firstRow", "bandRow"):
            tbl_pr.set(attr, "0")
        style_id = tbl_pr.find(qn("a:tableStyleId"))
        if style_id is not None:
            style_id.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # No Style, No Grid
        total = sum(widths)
        for i, weight in enumerate(widths):
            table.columns[i].width = int(w * weight / total)
        for r, row in enumerate(rows):
            table.rows[r].height = Inches(row_h)
            for c, value in enumerate(row):
                cell = table.cell(r, c)
                cell.margin_left = cell.margin_right = Inches(0.08)
                cell.margin_top = cell.margin_bottom = Inches(0.03)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.fill.solid()
                header = r == 0
                cell.fill.fore_color.rgb = _rgb(self.p.accent if header else (self.p.card if r % 2 else self.p.page))
                para = cell.text_frame.paragraphs[0]
                para.alignment = PP_ALIGN.RIGHT if c >= align_right_from else PP_ALIGN.LEFT
                run = para.add_run()
                run.text = value
                run.font.name = FONT
                run.font.size = Pt(10.5 if not header else 10)
                run.font.bold = header
                tone = self.p.on_accent if header else self.p.text
                if not header and c in (3, 4) and value[:1] in "+-":
                    tone = self.p.good if value[0] == "+" else self.p.bad
                run.font.color.rgb = _rgb(tone)
        return shape


# ---------------------------------------------------------------------------
# Content
# ---------------------------------------------------------------------------


def _agent_pipeline_names(analysis: dict) -> list[str]:
    # An agent can run more than once (e.g. a second analyst pass); list each once.
    return list(dict.fromkeys(step["name"] for step in analysis["agent_timeline"]))


def _merge_timeline(timeline: list[dict]) -> list[dict]:
    merged: dict[str, int] = {}
    for step in timeline:
        merged[step["name"]] = merged.get(step["name"], 0) + step["duration_ms"]
    return [{"name": n, "duration_ms": d} for n, d in merged.items()]


def _bridge_chart(deck: _Deck, slide, x, y, w, h, waterfall: list[dict]):
    """Waterfall as a stacked column. Under each driver bar sits the revenue
    carried forward from the steps before it, in a soft tint, so no bar
    floats over empty space. Driver bars carry their share of opening
    revenue in the axis label."""
    cats, base, up, down, total = [], [], [], [], []
    running = 0.0
    opening = waterfall[0]["value"] if waterfall else 0
    for step in waterfall:
        name = _clip(step["name"], 22)
        if step["type"] not in ("start", "end") and opening:
            name += f" ({_pct(step['value'] / abs(opening) * 100, True)})"
        cats.append(name)
        value = step["value"]
        if step["type"] in ("start", "end"):
            base.append(None); up.append(None); down.append(None); total.append(value)
            running = value
        elif value >= 0:
            base.append(running); up.append(value); down.append(None); total.append(None)
            running += value
        else:
            running += value
            base.append(running); up.append(None); down.append(-value); total.append(None)
    data = CategoryChartData()
    data.categories = cats
    for name, values in (("Base", base), ("Total", total), ("Increase", up), ("Decrease", down)):
        data.add_series(name, values)
    frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, x, y, w, h, data)
    chart = frame.chart
    deck.style_chart(chart)
    deck.style_axes(chart)
    plot = chart.plots[0]
    plot.gap_width = 45
    plot.overlap = 100
    fills = {"Base": deck.p.carry, "Total": deck.p.accent, "Increase": deck.p.good, "Decrease": deck.p.bad}
    for series in plot.series:
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = _rgb(fills[series.name])
        if series.name == "Base":
            continue  # its value is the previous bar's, so no label
        labels = series.data_labels
        labels.show_value = True
        labels.number_format = CHART_K
        labels.number_format_is_linked = False
        labels.font.size = Pt(10)
        labels.font.bold = True
        labels.font.color.rgb = _rgb(deck.p.on_accent if series.name == "Total" and deck.p.key == "finance"
                                     else deck.p.text)
        labels.position = XL_LABEL_POSITION.INSIDE_END if series.name == "Total" else XL_LABEL_POSITION.CENTER
        if series.name != "Total":
            labels.font.color.rgb = _rgb("FFFFFF")
    # Keep the axis from starting at zero when the bars sit near the top, so
    # the movements are visible instead of tiny slivers on a huge column.
    floor = min(v for v in base + total if v is not None)
    ceiling = max((b or 0) + (u or 0) + (t or 0) for b, u, t in zip(base, up, total))
    if floor > 0 and (ceiling - floor) < ceiling * 0.5:
        chart.value_axis.minimum_scale = max(0.0, round(floor * 0.8, -3))
    chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW


# A key driver is one of the fewest drivers that together explain this share
# of the total movement (Pareto); everything after that is non-key.
KEY_SHARE = 0.8


def _growth(d: dict | None) -> str:
    """A driver's change against its own prior-period total."""
    if not d:
        return ""
    return "new" if not d["prior"] else _pct(d["share_of_change_pct"], True)


def _pareto(rows: list[dict], value_key: str) -> list[dict]:
    """Sort by size of movement and add share (of total movement) and key."""
    ordered = sorted(rows, key=lambda r: -abs(r[value_key]))
    total = sum(abs(r[value_key]) for r in ordered) or 1.0
    covered, out = 0.0, []
    for r in ordered:
        share = abs(r[value_key]) / total
        out.append({**r, "share": share, "key": covered < KEY_SHARE})
        covered += share
    return out


def _classify_drivers(drivers: list[dict], transactions: list) -> list[dict]:
    """Drivers with kind (revenue or cost), favourable (revenue up / cost down),
    share of total movement and key flag."""
    kinds = {t.account_name: t.account_type.value for t in transactions}
    rows = []
    for d in drivers:
        kind = "revenue" if kinds.get(d["account"]) == "revenue" else "cost"
        rows.append({**d, "kind": kind, "favourable": (d["amount"] >= 0) == (kind == "revenue")})
    return _pareto(rows, "amount")


def _driver_chart(deck: _Deck, slide, x, y, w, h, rows: list[dict]):
    """Change per driver. Colour says favourable or not, a star marks key
    drivers, and each label pairs the amount with the driver's own % change."""
    top = rows[:8]
    data = CategoryChartData()
    data.categories = [("★ " if r["key"] else "") + _clip(r["driver"], 24) for r in top]
    data.add_series("Change", [r["amount"] for r in top])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, x, y, w, h, data).chart
    deck.style_chart(chart)
    deck.style_axes(chart)
    chart.category_axis.reverse_order = True
    chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    chart.category_axis.tick_labels.font.size = Pt(11)
    chart.value_axis.visible = False
    chart.value_axis.has_major_gridlines = False
    plot = chart.plots[0]
    plot.gap_width = 45
    series = plot.series[0]
    series.invert_if_negative = False
    for index, r in enumerate(top):
        point = series.points[index]
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = _rgb(deck.p.good if r["favourable"] else deck.p.bad)
        label = point.data_label
        label.position = XL_LABEL_POSITION.OUTSIDE_END
        run = label.text_frame.paragraphs[0].add_run()
        run.text = f"{_money(r['amount'], True)}  |  {_growth(r)}"
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = _rgb(deck.p.text)
    # Outside labels clip at the plot edge; leave room for them.
    high = max((r["amount"] for r in top), default=0)
    low = min((r["amount"] for r in top), default=0)
    if high > 0:
        chart.value_axis.maximum_scale = high * 1.55
    if low < 0:
        chart.value_axis.minimum_scale = low * 1.55


def _movement_pie(deck: _Deck, slide, x, y, w, h, rows: list[dict]):
    """Each key driver's share of the total movement, non-key drivers as one
    grey slice. Sizes are absolute so a decrease gets a slice too."""
    keys = [r for r in rows if r["key"]]
    rest = [r for r in rows if not r["key"]]
    slices = [(_clip(r["driver"], 26), r["share"], deck.p.slices[i % len(deck.p.slices)])
              for i, r in enumerate(keys)]
    if rest:
        slices.append((f"Non-key drivers ({len(rest)})", sum(r["share"] for r in rest), deck.p.muted))
    data = CategoryChartData()
    data.categories = [name for name, _, _ in slices]
    data.add_series("Share of movement", [share for _, share, _ in slices])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.PIE, x, y, w, h, data).chart
    deck.style_chart(chart, legend=True)
    chart.legend.position = XL_LEGEND_POSITION.RIGHT
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(10)
    chart.legend.font.color.rgb = _rgb(deck.p.text)
    series = chart.plots[0].series[0]
    for index, (_, _, colour) in enumerate(slices):
        point = series.points[index]
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = _rgb(colour)
        point.format.line.color.rgb = _rgb(deck.p.card)
    chart.plots[0].has_data_labels = True
    labels = chart.plots[0].data_labels
    labels.show_percentage = True
    labels.show_value = False
    labels.show_category_name = False
    labels.number_format = "0%"
    labels.number_format_is_linked = False
    labels.position = XL_LABEL_POSITION.CENTER
    labels.font.size = Pt(11)
    labels.font.bold = True
    labels.font.color.rgb = _rgb("FFFFFF")


def _recommendations(rows: list[dict], revenue_growth: float, revenue_rows: list[dict]) -> list[str]:
    """Actions that follow from the numbers; nothing here is model-written."""
    recs = []
    lead = next((r for r in revenue_rows if r["value"] > 0), None)
    if lead:
        recs.append(f"Scale {_clip(lead['name'], 30)}: {lead['share']:.0%} of the revenue movement "
                    f"({_money(lead['value'], True)}). Confirm it repeats next period before planning on it.")
    drag = next((r for r in revenue_rows if r["value"] < 0), None)
    if drag:
        recs.append(f"Recover {_clip(drag['name'], 30)} ({_money(drag['value'], True)}): review the decline "
                    f"with the account owner and set a recovery target.")
    cost = next((r for r in rows if r["kind"] == "cost" and r["amount"] > 0), None)
    if cost and cost["prior"] and _pace(cost["share_of_change_pct"] / 100, revenue_growth) == "faster than":
        recs.append(f"Contain {_clip(cost['driver'], 30)}: up {_growth(cost)}, faster than revenue "
                    f"({_pct(revenue_growth * 100, True)}). Review unit costs and contract terms.")
    elif cost:
        recs.append(f"Hold {_clip(cost['driver'], 30)} ({_growth(cost)}) at or below revenue growth to protect "
                    "margin as volume rises.")
    if lead and lead["share"] > 0.6:
        recs.append(f"Reduce concentration: one driver carries {lead['share']:.0%} of the movement; "
                    "build the next growth channel.")
    keys = [r for r in revenue_rows if r["key"]]
    if keys:
        recs.append(f"Track the {len(keys)} key revenue driver{'s' if len(keys) != 1 else ''} monthly; together "
                    f"they carry {sum(r['share'] for r in keys):.0%} of the movement.")
    if any(r["confidence"] < 0.8 for r in rows):
        recs.append("Validate the low-confidence drivers with account owners before external use.")
    else:
        recs.append("Sign off the key drivers with account owners, then share externally.")
    return recs[:4]


def _pace(rate: float, revenue_growth: float) -> str:
    """How a growth rate compares with revenue's, with half a point of slack."""
    if abs(rate - revenue_growth) < 0.005:
        return "in line with"
    return "faster than" if rate > revenue_growth else "slower than"


def _derived_risks(rows: list[dict], revenue_rows: list[dict], revenue_growth: float) -> list[str]:
    """Risks that follow from the driver figures, alongside the model's caveats."""
    risks = []
    lead = next((r for r in revenue_rows if r["value"] > 0), None)
    if lead and lead["share"] >= 0.5:
        risks.append(f"Concentration: {_clip(lead['name'], 30)} alone is {lead['share']:.0%} of the movement.")
    new = [r["driver"] for r in rows if r["kind"] == "revenue" and not r["prior"] and r["amount"] > 0]
    if new:
        risks.append(f"{', '.join(_clip(n, 24) for n in new[:2])} {'is' if len(new) == 1 else 'are'} new this period; "
                     "one period of revenue is not yet a trend.")
    cost = [r for r in rows if r["kind"] == "cost" and r["amount"] > 0 and r["prior"]]
    if cost:
        c = cost[0]
        risks.append(f"Margin: {_clip(c['driver'], 24)} costs rose {_growth(c)}, "
                     f"{_pace(c['share_of_change_pct'] / 100, revenue_growth)} revenue.")
    low = [r["driver"] for r in rows if r["confidence"] < 0.8]
    if low:
        risks.append(f"Lower confidence on {', '.join(_clip(n, 20) for n in low[:2])}; check the evidence first.")
    return risks


def _agent_chart(deck: _Deck, slide, x, y, w, h, timeline: list[dict]):
    data = CategoryChartData()
    data.categories = [_clip(s["name"], 26) for s in timeline]
    data.add_series("Duration (ms)", [s["duration_ms"] for s in timeline])
    frame = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, x, y, w, h, data)
    chart = frame.chart
    deck.style_chart(chart)
    deck.style_axes(chart, value_format='#,##0" ms"')
    chart.category_axis.reverse_order = True
    chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    chart.value_axis.visible = False
    chart.value_axis.has_major_gridlines = False
    chart.value_axis.maximum_scale = max(s["duration_ms"] for s in timeline) * 1.25
    plot = chart.plots[0]
    plot.gap_width = 50
    series = plot.series[0]
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = _rgb(deck.p.accent_2)
    labels = series.data_labels
    labels.show_value = True
    labels.number_format = '#,##0" ms"'
    labels.number_format_is_linked = False
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    labels.font.size = Pt(9)
    labels.font.color.rgb = _rgb(deck.p.text)


def build_pitch_deck(state: RunState, theme: str = "finance") -> bytes:
    palette = PALETTES.get(theme, PALETTES["finance"])
    deck = _Deck(palette)
    a = to_frontend_analysis(state)
    p = palette

    prior, current = a["periods"]["prior"], a["periods"]["current"]
    change = current["revenue"] - prior["revenue"]
    change_pct = (change / prior["revenue"] * 100) if prior["revenue"] else 0.0
    drivers = a["drivers"]
    steps = [s for s in a["waterfall"] if s["type"] in ("increase", "decrease")]
    gains = sorted((s for s in steps if s["value"] > 0), key=lambda s: -s["value"])
    losses = sorted((s for s in steps if s["value"] < 0), key=lambda s: s["value"])
    company = state.get("company_id") or "the business"
    profile = state.get("company_profile_core")
    if profile is not None and getattr(profile, "company_name", None):
        company = profile.company_name
    run_label = f"Ledgerlight  |  {company}  |  {prior['label']} to {current['label']}  |  Run {a['run_id']}"
    tone = p.good if change >= 0 else p.bad
    n_txn = len(state.get("transactions", []))
    n_facts = len(state.get("facts", []))
    n_evidence = sum(len(f.evidence_txn_ids) for f in state.get("facts", []))
    timeline = _merge_timeline(a["agent_timeline"])
    verified = a["grounding_badge"] or "n/a"
    risks = list(dict.fromkeys(a["risks_or_caveats"]))  # the model can repeat a caveat
    follow_ups = a["follow_up_questions"]
    growth = change / prior["revenue"] if prior["revenue"] else 0.0
    rows = _classify_drivers(drivers, state.get("transactions", []))
    by_name = {d["driver"]: d for d in drivers}
    revenue_rows = _pareto([{**x, "growth": _growth(by_name.get(x["name"]))} for x in steps], "value")
    rev_keys = [r for r in revenue_rows if r["key"]]
    up_total = sum(x["value"] for x in gains)
    down_total = sum(x["value"] for x in losses)
    keys = [r for r in rows if r["key"]]
    key_cover = sum(r["share"] for r in keys)
    cost_d = [r for r in rows if r["kind"] == "cost"]
    cost_change = sum(r["amount"] for r in cost_d)
    cost_prior = sum(r["prior"] for r in cost_d)
    cost_growth = cost_change / cost_prior if cost_prior else 0.0

    # ---- 1. Title -------------------------------------------------------
    s = deck.prs.slides.add_slide(deck.blank)
    deck.count += 1
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = _rgb(p.accent if p.key == "finance" else p.page)
    on_dark = p.key == "finance"
    ink = "FFFFFF" if on_dark else p.text
    soft = "C9D6EA" if on_dark else p.muted
    deck.rect(s, 0, 0, Inches(0.18), SLIDE_H, p.accent_2)
    deck.text(s, MARGIN, Inches(0.7), Inches(9), Inches(0.3), "LEDGERLIGHT  |  REVENUE VARIANCE REVIEW", size=12,
              bold=True, color=p.accent_2 if not on_dark else "9CC3E8")
    deck.text(s, MARGIN, Inches(1.25), Inches(11.5), Inches(1.5),
              f"{company}: revenue {'grew' if change >= 0 else 'fell'} {_money(abs(change))} "
              f"({_pct(abs(change_pct))}) from {prior['label']} to {current['label']}",
              size=38, bold=True, color=ink, line_spacing=0.95)
    deck.text(s, MARGIN, Inches(3.0), Inches(10.5), Inches(1.0), _first_sentence(a["headline"], 260),
              size=18, color=soft, line_spacing=1.05)
    kpis = [
        (_money(prior["revenue"]), f"{prior['label']} revenue", None),
        (_money(current["revenue"]), f"{current['label']} revenue", None),
        (_money(change, True), "Net change", tone),
        (f"{a['data_quality_score'] * 100:.0f}%", "Data quality", None),
    ]
    kw, gap = Inches(2.85), Inches(0.2)
    for i, (value, label, k_tone) in enumerate(kpis):
        x = MARGIN + i * (kw + gap)
        deck.rect(s, x, Inches(4.7), kw, Inches(1.35), p.card if not on_dark else "FFFFFF", p.border)
        deck.rect(s, x, Inches(4.7), Inches(0.06), Inches(1.35), k_tone or p.accent_2)
        deck.text(s, x + Inches(0.25), Inches(4.85), kw - Inches(0.4), Inches(0.3), label.upper(), size=10,
                  bold=True, color=p.muted if not on_dark else "5B6778")
        deck.text(s, x + Inches(0.25), Inches(5.2), kw - Inches(0.4), Inches(0.7), value, size=30, bold=True,
                  color=k_tone or (p.text if not on_dark else p.accent))
    deck.text(s, MARGIN, Inches(6.6), Inches(11), Inches(0.3),
              f"Prepared for leadership and partner teams  |  Run {a['run_id']}  |  Every figure traced to source transactions",
              size=11, color=soft)
    deck.notes(s, f"Opening. {a['headline']} Revenue moved from {_money(prior['revenue'])} to "
                  f"{_money(current['revenue'])}. The rest of the deck explains why, how sure we are, and what we do next.")

    # ---- 2. Problem and approach ---------------------------------------
    s = deck.slide("Context", "Explaining a revenue swing takes days and is hard to audit; we made it traceable",
                   run_label)
    col_w = Inches(5.9)
    deck.panel_title(s, MARGIN, Inches(1.85), col_w, "The problem")
    deck.text(s, MARGIN, Inches(2.3), col_w, Inches(2.6), [
        ("Finance teams reconcile month-on-month revenue by hand across ledgers, exports and spreadsheets.",
         {"bullet": True, "space_after": 9}),
        ("Narratives drift from the numbers: a figure quoted in a board note is rarely linked to the entries behind it.",
         {"bullet": True, "space_after": 9}),
        ("Leaders need three answers fast: what moved, why, and how much to trust the explanation.",
         {"bullet": True, "space_after": 9}),
    ], size=14, line_spacing=1.05)

    deck.card(s, MARGIN, Inches(4.95), col_w, Inches(1.8))
    deck.text(s, MARGIN + Inches(0.25), Inches(5.08), col_w, Inches(0.3), "SCOPE OF THIS RUN", size=10, bold=True,
              color=p.muted)
    stats = [(f"{n_txn:,}", "transactions"), (f"{n_facts}", "computed facts"), (f"{n_evidence}", "evidence links")]
    sw = (col_w - Inches(0.5)) / 3
    for i, (value, label) in enumerate(stats):
        deck.text(s, MARGIN + Inches(0.25) + i * sw, Inches(5.45), sw, Inches(0.55), value, size=28, bold=True,
                  color=p.accent if p.key == "finance" else p.accent)
        deck.text(s, MARGIN + Inches(0.25) + i * sw, Inches(6.02), sw, Inches(0.3), label, size=11, color=p.muted)
    deck.text(s, MARGIN + Inches(0.25), Inches(6.38), col_w - Inches(0.4), Inches(0.3),
              f"{prior['label']} vs {current['label']}  |  {company}", size=10, color=p.muted)

    rx = MARGIN + col_w + Inches(0.5)
    rw = SLIDE_W - MARGIN - rx
    deck.panel_title(s, rx, Inches(1.85), rw, "Our approach: a pipeline of specialised agents")
    names = _agent_pipeline_names(a) or ["Ingest", "Quality check", "Analyse", "Explain", "Verify", "Report"]
    names = names[:8]
    row_h = min(Inches(0.52), int((Inches(4.45)) / max(len(names), 1)))
    for i, name in enumerate(names):
        y = Inches(2.35) + i * row_h
        deck.rect(s, rx, y, Inches(0.36), Inches(0.36), p.accent, shape=MSO_SHAPE.OVAL)
        deck.text(s, rx, y + Inches(0.045), Inches(0.36), Inches(0.3), str(i + 1), size=11, bold=True,
                  color=p.on_accent, align=PP_ALIGN.CENTER)
        deck.text(s, rx + Inches(0.55), y + Inches(0.04), rw - Inches(0.6), Inches(0.3), name, size=14,
                  color=p.text)
        if i < len(names) - 1:
            deck.rect(s, rx + Inches(0.175), y + Inches(0.38), Emu(12700), row_h - Inches(0.4), p.border)
    deck.text(s, rx, Inches(6.55), rw, Inches(0.4),
              "Deterministic maths computes every figure; the language model only explains them.", size=11,
              color=p.muted, line_spacing=1.0)
    deck.notes(s, "Problem and approach. Stress that numbers are computed in code, not generated by the model, and "
                  "every fact keeps a pointer to the transactions behind it.")

    # ---- 3. Revenue bridge ---------------------------------------------
    best = gains[0] if gains else None
    worst = losses[0] if losses else None
    title3 = (f"Revenue moved {_money(change, True)} ({_pct(change_pct, True)}); "
              f"{_clip(best['name'], 30)} was the largest lift" if best else
              f"Revenue moved {_money(change, True)} ({_pct(change_pct, True)}) between the two periods")
    s = deck.slide("Analysis  |  Revenue bridge", title3, run_label)
    deck.card(s, MARGIN, Inches(1.8), Inches(7.4), Inches(5.0))
    _bridge_chart(deck, s, MARGIN + Inches(0.1), Inches(1.9), Inches(7.2), Inches(4.5), a["waterfall"])
    deck.text(s, MARGIN + Inches(0.25), Inches(6.47), Inches(7.0), Inches(0.25),
              "Grey is revenue carried forward; green adds, red reduces. % is against opening revenue.",
              size=10, color=p.muted)
    rx = MARGIN + Inches(7.7)
    rw = SLIDE_W - MARGIN - rx
    deck.panel_title(s, rx, Inches(1.8), rw, "Key drivers: increases vs decreases")
    tile_w = (rw - Inches(0.15)) / 2
    for i, (label, value, n, k_tone) in enumerate((("Increases", up_total, len(gains), p.good),
                                                  ("Decreases", down_total, len(losses), p.bad))):
        x = rx + i * (tile_w + Inches(0.15))
        deck.card(s, x, Inches(2.3), tile_w, Inches(1.0))
        deck.rect(s, x, Inches(2.3), Inches(0.06), Inches(1.0), k_tone)
        deck.text(s, x + Inches(0.2), Inches(2.38), tile_w - Inches(0.3), Inches(0.25), f"{label.upper()} ({n})",
                  size=9, bold=True, color=p.muted)
        deck.text(s, x + Inches(0.2), Inches(2.6), tile_w - Inches(0.3), Inches(0.4),
                  _money(value, True) if value else "$0", size=20, bold=True, color=k_tone if value else p.muted)
        if prior["revenue"]:
            deck.text(s, x + Inches(0.2), Inches(3.0), tile_w - Inches(0.3), Inches(0.25),
                      f"{_pct(value / prior['revenue'] * 100, True)} of opening revenue", size=9, color=p.muted)
    shown = revenue_rows[:4]
    if len(revenue_rows) > 4:
        rest = revenue_rows[3:]
        shown = revenue_rows[:3] + [{"name": f"{len(rest)} other drivers", "value": sum(r["value"] for r in rest),
                                      "share": sum(r["share"] for r in rest), "key": False, "growth": ""}]
    bar_w = rw - Inches(1.45)
    chip_w = Inches(0.8)
    for i, r in enumerate(shown):
        y = Inches(3.5) + i * Inches(0.56)
        r_tone = p.good if r["value"] >= 0 else p.bad
        deck.text(s, rx, y, rw - chip_w - Inches(1.05), Inches(0.26), _clip(r["name"], 24), size=12, bold=True)
        chip_x = rx + rw - chip_w - Inches(0.95)
        deck.rect(s, chip_x, y + Inches(0.02), chip_w, Inches(0.22), p.accent if r["key"] else None,
                  None if r["key"] else p.border, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        deck.text(s, chip_x, y + Inches(0.045), chip_w, Inches(0.2), "KEY" if r["key"] else "NON-KEY", size=8,
                  bold=True, color=p.on_accent if r["key"] else p.muted, align=PP_ALIGN.CENTER)
        deck.text(s, rx + rw - Inches(0.9), y, Inches(0.9), Inches(0.26), _money(r["value"], True), size=12,
                  bold=True, color=r_tone, align=PP_ALIGN.RIGHT)
        deck.rect(s, rx, y + Inches(0.34), bar_w, Inches(0.08), p.border)
        deck.rect(s, rx, y + Inches(0.34), max(int(bar_w * r["share"]), Emu(12700)), Inches(0.08),
                  r_tone if r["key"] else p.muted)
        detail = f"{r['share']:.0%}" + (f"  |  {r['growth']}" if r["growth"] else "")
        deck.text(s, rx + bar_w + Inches(0.1), y + Inches(0.28), Inches(1.35), Inches(0.2), detail, size=9,
                  color=p.muted, align=PP_ALIGN.RIGHT)
    deck.card(s, rx, Inches(5.8), rw, Inches(1.0))
    deck.text(s, rx + Inches(0.2), Inches(5.88), rw - Inches(0.4), Inches(0.25),
              "SO WHAT  (bar = share of movement  |  change vs own prior)", size=9, bold=True, color=p.muted)
    non_key = len(revenue_rows) - len(rev_keys)
    sowhat = (f"{len(rev_keys)} key driver{'s' if len(rev_keys) != 1 else ''} explain "
              f"{sum(r['share'] for r in rev_keys):.0%} of the movement; "
              + (f"{non_key} non-key make up the rest." if non_key else "there are no non-key drivers.")
              if rev_keys else f"Net effect {_money(change, True)}; no single driver dominates.")
    deck.text(s, rx + Inches(0.2), Inches(6.15), rw - Inches(0.4), Inches(0.6), sowhat, size=12.5, bold=True,
              color=tone, line_spacing=1.0)
    other = next((x for x in steps if x["name"] == "Other"), None)
    deck.notes(s, "Walk the chart left to right: opening revenue, each driver's contribution, closing revenue. "
                  "Green bars add revenue, red bars subtract; they tie exactly to both endpoints. On the right, "
                  "increases are compared with decreases, then each driver is ranked by its share of the movement. "
                  f"Key drivers are the fewest that together explain {KEY_SHARE:.0%} of it. "
                  + (f"Unattributed remainder is shown as 'Other' ({_money(other['value'], True)}), never hidden."
                     if other else "Drivers fully explain the movement; no unattributed remainder."))

    # ---- 4. Drivers -----------------------------------------------------
    title4 = (f"{len(keys)} of {len(rows)} drivers explain {key_cover:.0%} of the movement"
              + ("; costs grew faster than revenue" if cost_d and cost_change > 0 and cost_growth > growth else "")
              if rows else "Driver analysis")
    s = deck.slide("Analysis  |  Drivers", title4, run_label)
    lw = Inches(6.9)
    deck.panel_title(s, MARGIN, Inches(1.8), lw, "Change by driver: favourable vs unfavourable")
    deck.card(s, MARGIN, Inches(2.25), lw, Inches(3.6))
    if rows:
        _driver_chart(deck, s, MARGIN + Inches(0.05), Inches(2.3), lw - Inches(0.1), Inches(3.15), rows)
    deck.text(s, MARGIN + Inches(0.2), Inches(5.5), lw - Inches(0.4), Inches(0.3),
              "\u2605 key driver. Green is favourable (revenue up or cost down), red unfavourable. "
              "% is change against the driver's own prior period.", size=9, color=p.muted)
    rx = MARGIN + lw + Inches(0.3)
    rw = SLIDE_W - MARGIN - rx
    deck.panel_title(s, rx, Inches(1.8), rw, "Share of total movement: key vs non-key")
    deck.card(s, rx, Inches(2.25), rw, Inches(3.6))
    if rows:
        _movement_pie(deck, s, rx + Inches(0.05), Inches(2.35), rw - Inches(0.1), Inches(3.4), rows)
    avg_conf = sum(r["confidence"] for r in rows) / len(rows) * 100 if rows else 0
    rev_d = [r for r in rows if r["kind"] == "revenue"]
    rev_change = sum(r["amount"] for r in rev_d)
    tiles = [
        ("Key drivers", f"{len(keys)} of {len(rows)}",
         f"{key_cover:.0%} of all movement; average confidence {avg_conf:.0f}%", p.accent_2),
        ("Revenue drivers", _money(rev_change, True),
         f"{sum(r['amount'] > 0 for r in rev_d)} up, {sum(r['amount'] < 0 for r in rev_d)} down",
         p.good if rev_change >= 0 else p.bad),
        ("Cost drivers", _money(cost_change, True) if cost_d else "None",
         (f"{_pct(cost_growth * 100, True)} vs prior; {'unfavourable' if cost_change > 0 else 'favourable'}"
          if cost_d else "no cost movement in the top drivers"), p.bad if cost_change > 0 else p.good),
    ]
    tile_w = (CONTENT_W - Inches(0.4)) / 3
    for i, (label, value, sub, k_tone) in enumerate(tiles):
        x = MARGIN + i * (tile_w + Inches(0.2))
        deck.card(s, x, Inches(6.0), tile_w, Inches(0.8))
        deck.rect(s, x, Inches(6.0), Inches(0.06), Inches(0.8), k_tone)
        deck.text(s, x + Inches(0.2), Inches(6.08), tile_w - Inches(0.3), Inches(0.22), label.upper(), size=9,
                  bold=True, color=p.muted)
        deck.text(s, x + Inches(0.2), Inches(6.3), Inches(1.4), Inches(0.42), value, size=19, bold=True, color=k_tone)
        deck.text(s, x + Inches(1.65), Inches(6.32), tile_w - Inches(1.8), Inches(0.42), sub, size=10,
                  line_spacing=1.0, anchor=MSO_ANCHOR.MIDDLE)
    deck.notes(s, "The chart compares every driver: bar length is the dollar change, colour is whether it helps "
                  "or hurts, and the label adds its % change against its own prior period. The pie shows how "
                  "concentrated the movement is: key drivers individually, the rest grouped. Revenue and cost "
                  "drivers are kept apart in the tiles so a cost increase is never read as good news.")

    # ---- 5. Trust and evidence ------------------------------------------
    s = deck.slide("Assurance", "Every number is verified against source data before it reaches a slide", run_label)
    kw, kh, kg = Inches(2.9), Inches(1.25), Inches(0.2)
    guard = (a["guardrail_status"] or "n/a").replace("_", " ").title()
    tiles = [
        (f"{a['data_quality_score'] * 100:.0f}%", "Data quality score", None),
        (f"{a['confidence'] * 100:.0f}%", "Overall confidence", None),
        (verified.replace(" numbers verified", ""), "Numbers verified", p.good),
        (guard, "Guardrail status", None),
    ]
    for i, (value, label, k_tone) in enumerate(tiles):
        deck.kpi(s, MARGIN + i * (kw + kg), Inches(1.85), kw, kh, value, label, k_tone,
                 size=26 if len(value) <= 8 else 17)
    deck.panel_title(s, MARGIN, Inches(3.35), Inches(6.5), "Agent run: time spent per stage")
    deck.card(s, MARGIN, Inches(3.8), Inches(7.0), Inches(3.0))
    if timeline:
        _agent_chart(deck, s, MARGIN + Inches(0.05), Inches(3.85), Inches(6.9), Inches(2.9), timeline)
    rx = MARGIN + Inches(7.3)
    rw = SLIDE_W - MARGIN - rx
    deck.panel_title(s, rx, Inches(3.35), rw, "What backs each claim")
    deck.text(s, rx, Inches(3.85), rw, Inches(3.0), [
        ("Figures are computed in code from the ledger, not written by the language model.",
         {"bullet": True, "space_after": 8}),
        (f"{n_facts} facts link to {n_evidence} source transactions, down to file and row.",
         {"bullet": True, "space_after": 8}),
        ("Quoted numbers in the narrative are checked against those facts; anything unverified is flagged.",
         {"bullet": True, "space_after": 8}),
        ("Guardrails screen the uploaded data for injected instructions and formula-style content.",
         {"bullet": True, "space_after": 8}),
    ], size=12.5, line_spacing=1.05)
    deck.notes(s, f"Assurance. {verified}. Data quality {a['data_quality_score'] * 100:.0f}%, overall confidence "
                  f"{a['confidence'] * 100:.0f}%. This is the answer to 'can I trust this?'.")

    # ---- 6. Conclusions, risks and recommendations ---------------------
    s = deck.slide("Conclusions", "What we found, what could change it, and what we recommend", run_label)
    findings = [f"Revenue {'rose' if change >= 0 else 'fell'} {_money(abs(change))} ({_pct(abs(change_pct))}) to "
                f"{_money(current['revenue'])}."]
    if rev_keys:
        names = ", ".join(_clip(r["name"], 24) for r in rev_keys[:3])
        findings.append(f"Key revenue driver{'s' if len(rev_keys) > 1 else ''}: {names}, "
                        f"{sum(r['share'] for r in rev_keys):.0%} of the movement.")
    findings.append(f"Largest drag: {_clip(worst['name'], 30)} at {_money(worst['value'], True)}." if worst else
                    f"No revenue driver declined; increases total {_money(up_total, True)}.")
    if cost_d:
        findings.append(f"Costs in the top drivers moved {_money(cost_change, True)} ({_pct(cost_growth * 100, True)}), "
                        f"{_pace(cost_growth, growth)} revenue ({_pct(growth * 100, True)}).")
    findings.append(f"Evidence-backed: {verified.lower()}, data quality {a['data_quality_score'] * 100:.0f}%.")
    # Figure-based risks first; the model's caveats fill the rest.
    all_risks = _derived_risks(rows, revenue_rows, growth) + [_clip(r, 150) for r in risks]
    columns = [
        ("Key findings", findings, p.accent),
        ("Risks and caveats", all_risks[:4] or ["No material caveats were raised for this run."], p.bad),
        ("Recommendations", _recommendations(rows, growth, revenue_rows), p.good),
    ]
    cw = (CONTENT_W - Inches(0.6)) / 3
    for i, (heading, items, bar) in enumerate(columns):
        x = MARGIN + i * (cw + Inches(0.3))
        deck.card(s, x, Inches(1.85), cw, Inches(4.15))
        deck.rect(s, x, Inches(1.85), cw, Inches(0.06), bar)
        deck.text(s, x + Inches(0.25), Inches(2.05), cw - Inches(0.5), Inches(0.35), heading, size=14, bold=True,
                  color=p.accent if p.key == "finance" else p.text)
        deck.text(s, x + Inches(0.25), Inches(2.55), cw - Inches(0.5), Inches(3.4),
                  [(t, {"bullet": True, "space_after": 8}) for t in items], size=11.5, line_spacing=1.0)
    board = a.get("board_update")
    deck.card(s, MARGIN, Inches(6.15), CONTENT_W, Inches(0.7))
    deck.rect(s, MARGIN, Inches(6.15), Inches(0.06), Inches(0.7), p.accent_2)
    deck.text(s, MARGIN + Inches(0.25), Inches(6.21), Inches(3), Inches(0.22), "BOARD UPDATE, ONE LINE", size=9,
              bold=True, color=p.muted)
    deck.text(s, MARGIN + Inches(0.25), Inches(6.42), CONTENT_W - Inches(0.5), Inches(0.4),
              _clip(board, 230) if board else _first_sentence(a["summary"], 230), size=11, line_spacing=1.0)
    deck.notes(s, (board or a["summary"] or "") + "\n\nRisks: " + "; ".join(risks)
               + "\n\nRecommendations are derived from the driver figures, not written by the model.")

    # ---- 7. Future work -------------------------------------------------
    s = deck.slide("Next steps", "From a one-off analysis to a monthly, always-on revenue review", run_label)
    cw = (CONTENT_W - Inches(0.6)) / 3
    columns = [
        ("Open questions from this run", [_clip(q, 105) for q in follow_ups[:4]] or
         ["Which drivers repeat month to month, and which were one-offs?"]),
        ("Product roadmap", [
            "Multi-period trends and seasonality, beyond a two-period comparison",
            "Forecasting and what-if scenarios on the top drivers",
            "Live connectors to the ledger and billing systems",
            "Scheduled monthly board pack, delivered automatically",
        ]),
        ("Asks of this audience", [
            "Confirm the drivers and caveats on slide 6 with account owners",
            "Nominate a pilot team and data source for the next close",
            "Agree the confidence threshold for external reporting",
            "Review the Power BI report for the same run",
        ]),
    ]
    for i, (heading, items) in enumerate(columns):
        x = MARGIN + i * (cw + Inches(0.3))
        deck.card(s, x, Inches(1.85), cw, Inches(4.2))
        deck.rect(s, x, Inches(1.85), cw, Inches(0.06), [p.accent, p.accent_2, p.good][i])
        deck.text(s, x + Inches(0.25), Inches(2.05), cw - Inches(0.5), Inches(0.35), heading, size=14, bold=True,
                  color=p.accent if p.key == "finance" else p.text)
        deck.text(s, x + Inches(0.25), Inches(2.6), cw - Inches(0.5), Inches(3.4),
                  [(t, {"bullet": True, "space_after": 9}) for t in items], size=12.5, line_spacing=1.0)
    deck.rect(s, MARGIN, Inches(6.25), CONTENT_W, Inches(0.55), p.accent)
    deck.text(s, MARGIN + Inches(0.3), Inches(6.25), CONTENT_W - Inches(0.6), Inches(0.55),
              "Next step: agree the pilot scope and run this review on the next period close.", size=14, bold=True,
              color=p.on_accent, anchor=MSO_ANCHOR.MIDDLE)
    deck.notes(s, "Close with the asks. The open questions on the left come straight from the analysis; "
                  "the roadmap is the proposal.")

    buffer = io.BytesIO()
    deck.prs.save(buffer)
    return buffer.getvalue()
