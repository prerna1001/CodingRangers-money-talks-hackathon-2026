from __future__ import annotations

import csv
import json
from io import BytesIO, StringIO

from .ocr import extract_text
from .types import ExtractionTable, SourceLocation


def _text_table(text: str, source: SourceLocation) -> ExtractionTable | None:
    lines = [line for line in text.splitlines() if line.strip()]
    for delimiter in (",", "\t", "|"):
        if len(lines) >= 2 and delimiter in lines[0]:
            reader = csv.DictReader(StringIO("\n".join(lines)), delimiter=delimiter)
            rows = [{str(k): str(v or "") for k, v in row.items()} for row in reader]
            if rows and reader.fieldnames:
                return ExtractionTable(list(reader.fieldnames), rows, source)
    return None


def extract_csv(raw: bytes) -> tuple[list[ExtractionTable], str, list[str]]:
    text = raw.decode("utf-8-sig")
    reader = csv.DictReader(StringIO(text))
    rows = [{str(k): str(v or "") for k, v in row.items()} for row in reader]
    headers = list(reader.fieldnames or [])
    return [ExtractionTable(headers, rows, SourceLocation(table=1))], text, []


def extract_json(raw: bytes) -> tuple[list[ExtractionTable], str, list[str]]:
    value = json.loads(raw.decode("utf-8"))
    rows = value if isinstance(value, list) else value.get("rows", []) if isinstance(value, dict) else []
    if rows and all(isinstance(row, dict) for row in rows):
        headers = list(dict.fromkeys(key for row in rows for key in row))
        return [ExtractionTable(headers, [{str(k): str(v or "") for k, v in row.items()} for row in rows], SourceLocation(table=1))], "", []
    return [], json.dumps(value), ["JSON contains no tabular row objects."]


def extract_xlsx(raw: bytes) -> tuple[list[ExtractionTable], str, list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("XLSX extraction requires openpyxl.") from exc
    workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    tables: list[ExtractionTable] = []
    for sheet in workbook.worksheets:
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            continue
        headers = [str(value or "").strip() for value in values[0]]
        if not any(headers):
            continue
        rows = []
        for row in values[1:]:
            if any(value not in (None, "") for value in row):
                rows.append({headers[i]: str(value or "") for i, value in enumerate(row) if headers[i]})
        tables.append(ExtractionTable(headers, rows, SourceLocation(sheet=sheet.title, table=1)))
    return tables, "", []


def extract_docx(raw: bytes) -> tuple[list[ExtractionTable], str, list[str]]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("DOCX extraction requires python-docx.") from exc
    document = Document(BytesIO(raw))
    tables: list[ExtractionTable] = []
    for index, table in enumerate(document.tables, start=1):
        matrix = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if len(matrix) < 2:
            continue
        headers = matrix[0]
        rows = [{headers[i]: value for i, value in enumerate(row) if i < len(headers) and headers[i]} for row in matrix[1:]]
        tables.append(ExtractionTable(headers, rows, SourceLocation(table=index)))
    return tables, "\n".join(paragraph.text for paragraph in document.paragraphs), []


def extract_pdf(raw: bytes) -> tuple[list[ExtractionTable], str, list[str], int, bool, float | None]:
    try:
        import fitz
    except ImportError as exc:
        raise RuntimeError("PDF extraction requires PyMuPDF.") from exc
    pdf = fitz.open(stream=raw, filetype="pdf")
    if len(pdf) > 100:
        raise ValueError("PDF exceeds the 100-page safety limit.")
    pages: list[str] = []
    tables: list[ExtractionTable] = []
    warnings: list[str] = []
    ocr_scores: list[float] = []
    used_ocr = False
    for page_index, page in enumerate(pdf, start=1):
        page_text = page.get_text("text").strip()
        # Native PDFs often retain the actual accounting table even when the
        # surrounding report is prose. Prefer it over heuristic text parsing.
        try:
            for table_index, found in enumerate(page.find_tables().tables, start=1):
                matrix = found.extract()
                if len(matrix) < 2:
                    continue
                headers = [str(value or "").strip() for value in matrix[0]]
                if not any(headers):
                    continue
                rows = [
                    {headers[i]: str(value or "") for i, value in enumerate(row) if i < len(headers) and headers[i]}
                    for row in matrix[1:]
                ]
                if rows:
                    tables.append(ExtractionTable(headers, rows, SourceLocation(page=page_index, table=table_index)))
        except (AttributeError, RuntimeError, ValueError):
            # Older PyMuPDF versions or layout-less pages still provide text.
            pass
        if len(page_text) < 30:
            used_ocr = True
            try:
                pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                result = extract_text(pixmap.tobytes("png"))
                page_text = result.text
                warnings.extend(result.warnings)
                if result.confidence is not None:
                    ocr_scores.append(result.confidence)
            except RuntimeError as exc:
                warnings.append(f"Page {page_index} needs OCR but could not be processed: {exc}")
        pages.append(page_text)
        if not tables:
            detected = _text_table(page_text, SourceLocation(page=page_index, table=1))
            if detected:
                detected.confidence = ocr_scores[-1] if ocr_scores else 1.0
                tables.append(detected)
    return tables, "\n\n".join(pages), warnings, len(pdf), used_ocr, sum(ocr_scores) / len(ocr_scores) if ocr_scores else None


def extract_image(raw: bytes) -> tuple[list[ExtractionTable], str, list[str], float | None]:
    result = extract_text(raw)
    table = _text_table(result.text, SourceLocation(page=1, table=1))
    if table is not None:
        table.confidence = result.confidence or 0.0
    return [table] if table else [], result.text, result.warnings, result.confidence
